#!/usr/bin/env python3
"""ANUBIS RAG — Ingestion Job (Kubernetes edition).

Ported from the .17 anubis_ingest_parallel_gpu.py. Same pipeline and guarantees:
  * parse TXT/PDF/DOCX/XLSX in parallel across CPU cores
  * fixed-window chunking (1400 chars / 200 overlap) — unchanged policy
  * SHA-256 file-hash dedup against Qdrant before re-ingesting
  * deterministic UUIDv5 point ids => idempotent re-ingestion
  * batched upsert

The ONE change: embeddings come from the shared embedder service (EMBED_URL),
not a torch model in this process. That keeps index-time and query-time vectors
identical and makes this image tiny (no torch/CUDA).
"""
import argparse
import hashlib
import os
import sys
import uuid
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Dict, Any

import requests

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".xlsx"}

QDRANT_URL = os.environ.get("QDRANT_URL", "http://qdrant:6333")
COLLECTION = os.environ.get("COLLECTION_NAME", "anubis_rag_chunks")
EMBED_URL = os.environ.get("EMBED_URL", "http://embedder:8080")


# ── Hashing ──
def file_hash(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


# ── Parsers (worker processes) ──
def _parse_txt(path):
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def _parse_pdf(path):
    import fitz
    out = []
    with fitz.open(path) as doc:
        for page in doc:
            t = page.get_text()
            if t:
                out.append(t)
    return "\n".join(out)


def _parse_docx(path):
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sheet: {sheet.title}]")
        for r in sheet.iter_rows(values_only=True):
            line = " | ".join(str(c) for c in r if c is not None)
            if line.strip():
                rows.append(line)
    return "\n".join(rows)


PARSERS = {".txt": _parse_txt, ".pdf": _parse_pdf, ".docx": _parse_docx, ".xlsx": _parse_xlsx}


def parse_file(path: str) -> Dict[str, Any]:
    p = Path(path)
    ext = p.suffix.lower()
    res = {"path": path, "name": p.name, "stem": p.stem, "ext": ext,
           "text": "", "error": None, "doc_hash": ""}
    try:
        res["doc_hash"] = file_hash(path)
        parser = PARSERS.get(ext)
        if parser is None:
            res["error"] = f"unsupported extension: {ext}"
            return res
        res["text"] = parser(path)
    except Exception as e:
        res["error"] = str(e)
    return res


# ── Chunking (unchanged policy) ──
def chunk_text(text: str, chunk_chars: int = 1400, overlap_chars: int = 200) -> List[str]:
    text = (text or "").strip()
    if not text:
        return []
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(start + chunk_chars, n)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(0, end - overlap_chars)
    return chunks


# ── Embedding via the shared service ──
def embed_texts(texts: List[str], batch_size: int = 64) -> List[List[float]]:
    vectors = []
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        r = requests.post(f"{EMBED_URL}/embed",
                          json={"texts": batch, "normalize": True, "batch_size": batch_size},
                          timeout=300)
        r.raise_for_status()
        vectors.extend(r.json()["vectors"])
    return vectors


def embed_dim() -> int:
    r = requests.get(f"{EMBED_URL}/health", timeout=30)
    r.raise_for_status()
    return int(r.json()["dim"])


# ── Qdrant helpers ──
def get_ingested_hashes(client, collection: str) -> set:
    hashes, offset = set(), None
    while True:
        results, offset = client.scroll(
            collection_name=collection, scroll_filter=None,
            limit=1000, offset=offset, with_payload=["doc_hash"])
        for pt in results:
            h = pt.payload.get("doc_hash")
            if h:
                hashes.add(h)
        if offset is None:
            break
    return hashes


def ensure_collection(client, collection: str, vector_size: int):
    from qdrant_client.http import models as qm
    if collection in {c.name for c in client.get_collections().collections}:
        return
    client.create_collection(
        collection_name=collection,
        vectors_config=qm.VectorParams(size=vector_size, distance=qm.Distance.COSINE),
        on_disk_payload=True,
    )


def batched_upsert(client, collection, ids, vectors, payloads, batch_size=128):
    from qdrant_client.http import models as qm
    for start in range(0, len(ids), batch_size):
        end = min(start + batch_size, len(ids))
        client.upsert(collection_name=collection,
                      points=qm.Batch(ids=ids[start:end], vectors=vectors[start:end],
                                      payloads=payloads[start:end]))


def main() -> int:
    ap = argparse.ArgumentParser(description="ANUBIS RAG — Ingestion Job")
    ap.add_argument("directory", help="Directory to ingest (recursive)")
    ap.add_argument("--chunk-chars", type=int, default=1400)
    ap.add_argument("--overlap-chars", type=int, default=200)
    ap.add_argument("--embed-batch", type=int, default=64)
    ap.add_argument("--upsert-batch", type=int, default=128)
    ap.add_argument("--workers", type=int, default=None)
    args = ap.parse_args()

    ingest_dir = Path(args.directory).expanduser().resolve()
    if not ingest_dir.is_dir():
        print(f"[ERROR] Not a directory: {ingest_dir}")
        return 2

    from qdrant_client import QdrantClient
    client = QdrantClient(url=QDRANT_URL)
    dim = embed_dim()
    ensure_collection(client, COLLECTION, dim)

    workers = args.workers or os.cpu_count() or 4
    print("=" * 60)
    print(" ANUBIS RAG — Ingestion Job")
    print(f"  Directory : {ingest_dir}")
    print(f"  Collection: {COLLECTION}")
    print(f"  Embedder  : {EMBED_URL} (dim {dim})")
    print(f"  Chunk     : {args.chunk_chars}/{args.overlap_chars}  workers {workers}")
    print("=" * 60)
    t0 = time.time()

    all_files = [str(f) for f in ingest_dir.rglob("*")
                 if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS]
    if not all_files:
        print("[WARN] No supported files found.")
        return 0
    print(f"[1/4] Discovered {len(all_files)} files")

    parsed = []
    with ProcessPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(parse_file, f): f for f in all_files}
        for fut in as_completed(futures):
            try:
                parsed.append(fut.result())
            except Exception as e:
                print(f"  [ERROR] {futures[fut]}: {e}")
    print(f"[2/4] Parsed {len(parsed)} files")

    existing = get_ingested_hashes(client, COLLECTION)
    print(f"[3/4] {len(existing)} existing doc hashes (dedup)")

    texts, meta = [], []
    stats = {"ingested": 0, "dedup": 0, "empty": 0, "errors": 0}
    for r in parsed:
        if r["error"]:
            stats["errors"] += 1
            print(f"  [ERROR] {r['name']}: {r['error']}")
            continue
        if r["doc_hash"] in existing:
            stats["dedup"] += 1
            continue
        chunks = chunk_text(r["text"], args.chunk_chars, args.overlap_chars)
        if not chunks:
            stats["empty"] += 1
            continue
        for idx, ch in enumerate(chunks):
            texts.append(ch)
            meta.append({"doc_hash": r["doc_hash"], "source_file": r["name"],
                         "source_path": r["path"], "stem": r["stem"], "chunk_index": idx})
        stats["ingested"] += 1

    if not texts:
        print("[INFO] No new chunks.")
        print(stats)
        return 0

    print(f"[4/4] Embedding {len(texts)} chunks via {EMBED_URL} ...")
    vectors = embed_texts(texts, batch_size=args.embed_batch)

    now = datetime.now(timezone.utc).isoformat()
    ids, payloads = [], []
    for i, m in enumerate(meta):
        ids.append(str(uuid.uuid5(uuid.NAMESPACE_URL, f"{m['doc_hash']}_{m['chunk_index']}")))
        payloads.append({
            "chunk_id": f"{m['stem']}_chunk_{m['chunk_index']}",
            "text": texts[i], "source_file": m["source_file"],
            "source_path": m["source_path"], "chunk_index": m["chunk_index"],
            "doc_hash": m["doc_hash"], "ingested_at": now,
        })
    batched_upsert(client, COLLECTION, ids, vectors, payloads, batch_size=args.upsert_batch)

    print(f"[DONE] upserted {len(ids)} chunks in {time.time() - t0:.1f}s  {stats}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
